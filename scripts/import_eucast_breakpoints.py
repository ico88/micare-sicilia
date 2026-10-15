"""Convert the official EUCAST Clinical Breakpoint Tables (.xlsx) into the
reference CSVs used by the app.

    python scripts/import_eucast_breakpoints.py path/to/v_16.1_Breakpoint_Tables.xlsx 16.1

Writes into app/resources/breakpoints/:
  eucast_v<ver>_rows.csv      every MIC row of the organism sheets we use, parsed
  eucast_v<ver>.csv           one chosen row per pathogen code x antibiotic code

Footnote markers in the EUCAST workbook are superscripts inside the same cell
("8" + superscript "4"). Read as plain text they merge into a wrong value
("84"), so cells are read as rich text and superscript runs are dropped.
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app.services.breakpoints import (  # noqa: E402
    ANTIBIOTIC_NAMES,
    PATHOGEN_GROUPS,
    RESOURCE_DIR,
)


def cell_text(value) -> tuple[str, str]:
    """(main text, footnote markers) of a cell, superscripts split off."""
    if value is None:
        return "", ""
    if isinstance(value, CellRichText):
        main, notes = [], []
        for block in value:
            if isinstance(block, TextBlock):
                sup = block.font is not None and block.font.vertAlign == "superscript"
                (notes if sup else main).append(block.text)
            else:
                main.append(str(block))
        return " ".join("".join(main).split()), ",".join(n.strip() for n in notes if n.strip())
    return " ".join(str(value).split()), ""


def parse_breakpoint(text: str) -> tuple[float | None, str, bool]:
    """'8' -> (8, 'ok', False); '(2)' -> (2, 'ok', True); 'IE' -> (None, 'IE', False)."""
    text = text.strip()
    bracketed = text.startswith("(") and text.endswith(")")
    core = text.strip("()").strip()
    try:
        return float(core), "ok", bracketed
    except ValueError:
        pass
    if core == "-":
        return None, "non_applicabile", bracketed
    if core.upper() == "IE":
        return None, "IE", bracketed
    if core.upper() == "NA":
        return None, "NA", bracketed
    if core.lower().startswith("note"):
        return None, "nota", bracketed
    return None, "non_interpretabile", bracketed


def read_rows(xlsx: Path) -> list[dict]:
    wb = openpyxl.load_workbook(xlsx, rich_text=True)
    sheets = sorted({group["sheet"] for group in PATHOGEN_GROUPS.values()})
    rows = []
    for sheet in sheets:
        ws = wb[sheet]
        in_mic_block = False
        for r in range(1, ws.max_row + 1):
            name, name_notes = cell_text(ws.cell(r, 1).value)
            s_text, s_notes = cell_text(ws.cell(r, 2).value)
            r_text, r_notes = cell_text(ws.cell(r, 3).value)
            if s_text == "S ≤" and r_text == "R >":
                in_mic_block = True  # header of a drug-class block
                continue
            if not in_mic_block or not name or not s_text:
                continue
            s_val, s_status, s_br = parse_breakpoint(s_text)
            r_val, r_status, r_br = parse_breakpoint(r_text)
            rows.append({
                "sheet": sheet,
                "excel_row": r,
                "agent": name,
                "agent_notes": name_notes,
                "s_text": s_text,
                "r_text": r_text,
                "s_le": s_val,
                "r_gt": r_val,
                "status": s_status if s_status != "ok" else r_status,
                "bracketed": s_br or r_br,
                "notes": ",".join(n for n in (s_notes, r_notes) if n),
            })
    return rows


_EXCLUDE = ("screen only", "(meningitis)", "(endocarditis and meningitis)")
# Lower = preferred. Surveillance convention: systemic / iv, non-meningitis indications.
_PENALTIES = [
    ("uncomplicated uti only", 3),
    ("infections originating from the urinary tract", 2),
    ("oral", 2),
    ("(pneumonia)", 1),
]


def choose(rows: list[dict], pathogen: str, antibiotic: str) -> dict | None:
    group = PATHOGEN_GROUPS[pathogen]
    name = ANTIBIOTIC_NAMES[antibiotic]
    pattern = re.compile(rf"^{re.escape(name)}(?=$|[\s,(])", re.IGNORECASE)
    candidates = []
    for row in rows:
        if row["sheet"] != group["sheet"] or not pattern.match(row["agent"]):
            continue
        text = row["agent"].lower()
        if any(word in text for word in _EXCLUDE):
            continue
        qualifier = row["agent"][len(name):]
        if not _species_ok(qualifier, group):
            continue
        penalty = sum(p for word, p in _PENALTIES if word in text)
        # A row naming the species beats a generic one
        specific = any(s.lower() in qualifier.lower() for s in group["species"])
        # A numeric breakpoint for the species wins over a "Note"/IE row; the chosen
        # indication is written to the CSV so a microbiologist can review it.
        numeric = 0 if row["status"] == "ok" else 1
        candidates.append((numeric, penalty, 0 if specific else 1, row["excel_row"], row))
    if not candidates:
        return None
    return sorted(candidates, key=lambda c: c[:4])[0][4]


def _species_ok(qualifier: str, group: dict) -> bool:
    """Skip rows restricted to other species ('Salmonella spp.', 'Coagulase-negative staphylococci')."""
    q = re.sub(r"\([^)]*\)", " ", qualifier.lower())  # drop indications in brackets
    if all(word in {"iv", "oral"} for word in q.replace(",", " ").split()):
        return True  # no species restriction
    if any(s.lower() in q for s in group["species"]):
        return True
    if "except" in q:
        return not any(s.lower() in q.split("except", 1)[1] for s in group["species"])
    return "other" in q and any(word in q for word in group["group_words"])


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("version", help="e.g. 16.1")
    args = parser.parse_args()

    rows = read_rows(args.xlsx)
    RESOURCE_DIR.mkdir(parents=True, exist_ok=True)
    rows_path = RESOURCE_DIR / f"eucast_v{args.version}_rows.csv"
    with rows_path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)

    chosen = []
    for pathogen in PATHOGEN_GROUPS:
        for antibiotic in ANTIBIOTIC_NAMES:
            row = choose(rows, pathogen, antibiotic)
            chosen.append({
                "pathogen": pathogen,
                "antibiotic": antibiotic,
                "sheet": row["sheet"] if row else PATHOGEN_GROUPS[pathogen]["sheet"],
                "excel_row": row["excel_row"] if row else "",
                "eucast_agent": row["agent"] if row else "",
                "s_le": row["s_le"] if row else "",
                "r_gt": row["r_gt"] if row else "",
                "status": row["status"] if row else "assente",
                "bracketed": row["bracketed"] if row else "",
                "notes": row["notes"] if row else "",
            })
    chosen_path = RESOURCE_DIR / f"eucast_v{args.version}.csv"
    with chosen_path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(chosen[0]))
        writer.writeheader()
        writer.writerows(chosen)
    usable = sum(1 for c in chosen if c["status"] == "ok")
    print(f"{len(rows)} righe MIC lette; {usable}/{len(chosen)} coppie con breakpoint numerico -> {chosen_path}")


if __name__ == "__main__":
    main()
